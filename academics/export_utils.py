# academics/export_utils.py
"""
Export utilities for generating PDF and Excel reports.
Supports CO-PO mapping reports, attainment reports, and other institutional documents.
"""
import io
from datetime import datetime
from django.http import HttpResponse
from django.template.loader import get_template
from django.conf import settings


def export_co_po_mapping_pdf(course, co_po_mappings):
    """
    Generate a PDF report for CO-PO mapping of a specific course.
    
    Args:
        course: Course object
        co_po_mappings: QuerySet or list of COPOMapping objects
        
    Returns:
        HttpResponse: PDF file response
    """
    try:
        from reportlab.lib.pagesizes import letter, A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
    except ImportError:
        raise ImportError("ReportLab is required for PDF export. Install with: pip install reportlab")
    
    # Create the HttpResponse object with PDF headers
    response = HttpResponse(content_type='application/pdf')
    filename = f"CO_PO_Mapping_{course.code}_{datetime.now().strftime('%Y%m%d')}.pdf"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    # Create the PDF object
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, 
                           rightMargin=0.5*inch, leftMargin=0.5*inch,
                           topMargin=0.75*inch, bottomMargin=0.5*inch)
    
    # Container for the 'Flowable' objects
    elements = []
    
    # Define styles
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=16,
        textColor=colors.HexColor('#604993'),
        spaceAfter=12,
        alignment=TA_CENTER,
        fontName='Helvetica-Bold'
    )
    
    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontSize=12,
        textColor=colors.HexColor('#604993'),
        spaceAfter=10,
        fontName='Helvetica-Bold'
    )
    
    normal_style = styles['Normal']
    
    # Title
    title = Paragraph(f"Course Outcome - Program Outcome Mapping", title_style)
    elements.append(title)
    elements.append(Spacer(1, 0.2*inch))
    
    # Course Information
    course_info = [
        ["Course Code:", course.code],
        ["Course Name:", course.name],
        ["Department:", course.department.name if course.department else "N/A"],
        ["Credits:", str(course.credits)],
        ["Semester:", course.semester.name if course.semester else "N/A"],
    ]
    
    course_table = Table(course_info, colWidths=[2*inch, 5*inch])
    course_table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#604993')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 12),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    
    elements.append(course_table)
    elements.append(Spacer(1, 0.3*inch))
    
    # CO-PO Mapping Matrix
    elements.append(Paragraph("CO-PO Mapping Matrix", heading_style))
    elements.append(Spacer(1, 0.1*inch))
    
    # Get unique COs and POs
    from collections import defaultdict
    co_dict = defaultdict(dict)
    pos = set()
    
    for mapping in co_po_mappings:
        co_code = mapping.course_outcome.code
        po_code = mapping.program_outcome.code
        level = mapping.correlation_level
        co_dict[co_code][po_code] = level
        pos.add(po_code)
    
    pos = sorted(pos, key=lambda x: (x[:2], int(x[2:]) if x[2:].isdigit() else 0))
    cos = sorted(co_dict.keys(), key=lambda x: (x[:2], int(x[2:]) if x[2:].isdigit() else 0))
    
    # Build table data
    table_data = [['CO \\ PO'] + pos]
    
    for co in cos:
        row = [co]
        for po in pos:
            level = co_dict[co].get(po, 0)
            if level == 0:
                row.append('-')
            elif level == 1:
                row.append('L')
            elif level == 2:
                row.append('M')
            elif level == 3:
                row.append('H')
        table_data.append(row)
    
    # Calculate column widths
    num_cols = len(pos) + 1
    col_width = 6.5*inch / num_cols
    
    mapping_table = Table(table_data, colWidths=[col_width] * num_cols)
    mapping_table.setStyle(TableStyle([
        # Header row
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#604993')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        # First column
        ('BACKGROUND', (0, 1), (0, -1), colors.HexColor('#F0F0F0')),
        ('FONTNAME', (0, 1), (0, -1), 'Helvetica-Bold'),
        # Grid
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (1, 1), (-1, -1), [colors.white, colors.HexColor('#F9F9F9')]),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    
    elements.append(mapping_table)
    elements.append(Spacer(1, 0.2*inch))
    
    # Legend
    legend_data = [
        ['Legend:', 'H = High (3)', 'M = Medium (2)', 'L = Low (1)', '- = None (0)']
    ]
    legend_table = Table(legend_data, colWidths=[1*inch, 1.3*inch, 1.5*inch, 1.2*inch, 1.3*inch])
    legend_table.setStyle(TableStyle([
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('FONTNAME', (0, 0), (0, 0), 'Helvetica-Bold'),
        ('TEXTCOLOR', (0, 0), (0, 0), colors.HexColor('#604993')),
        ('LEFTPADDING', (0, 0), (-1, -1), 2),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    
    elements.append(legend_table)
    elements.append(Spacer(1, 0.3*inch))
    
    # Footer
    footer_text = f"Generated on: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
    footer = Paragraph(footer_text, ParagraphStyle(
        'Footer',
        parent=styles['Normal'],
        fontSize=8,
        textColor=colors.grey,
        alignment=TA_CENTER
    ))
    elements.append(footer)
    
    # Build PDF
    doc.build(elements)
    
    # Get the value of the BytesIO buffer and write it to the response
    pdf = buffer.getvalue()
    buffer.close()
    response.write(pdf)
    
    return response


def export_co_attainment_pdf(co_attainments, filters=None):
    """
    Generate a PDF report for Course Outcome attainment.
    
    Args:
        co_attainments: QuerySet of CourseOutcomeAttainment objects
        filters: Dictionary of applied filters (optional)
        
    Returns:
        HttpResponse: PDF file response
    """
    try:
        from reportlab.lib.pagesizes import letter, A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.enums import TA_CENTER
    except ImportError:
        raise ImportError("ReportLab is required for PDF export. Install with: pip install reportlab")
    
    response = HttpResponse(content_type='application/pdf')
    filename = f"CO_Attainment_Report_{datetime.now().strftime('%Y%m%d')}.pdf"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                           rightMargin=0.5*inch, leftMargin=0.5*inch,
                           topMargin=0.75*inch, bottomMargin=0.5*inch)
    
    elements = []
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=16,
        textColor=colors.HexColor('#604993'),
        spaceAfter=12,
        alignment=TA_CENTER,
        fontName='Helvetica-Bold'
    )
    
    title = Paragraph("Course Outcome Attainment Report", title_style)
    elements.append(title)
    elements.append(Spacer(1, 0.3*inch))
    
    # Build table data
    table_data = [['Course Code', 'Course Name', 'CO Code', 'Academic Year', 'Attainment %']]
    
    for attainment in co_attainments:
        row = [
            attainment.course_outcome.course.code,
            Paragraph(attainment.course_outcome.course.name[:40], styles['Normal']),
            attainment.course_outcome.code,
            f"{attainment.academic_year.start_date.year}-{attainment.academic_year.end_date.year}",
            f"{attainment.attainment_percentage:.2f}%" if attainment.attainment_percentage else "N/A"
        ]
        table_data.append(row)
    
    attainment_table = Table(table_data, colWidths=[1*inch, 2.5*inch, 0.8*inch, 1.2*inch, 1*inch])
    attainment_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#604993')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9F9F9')]),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    
    elements.append(attainment_table)
    elements.append(Spacer(1, 0.3*inch))
    
    footer_text = f"Generated on: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
    footer = Paragraph(footer_text, ParagraphStyle(
        'Footer',
        parent=styles['Normal'],
        fontSize=8,
        textColor=colors.grey,
        alignment=TA_CENTER
    ))
    elements.append(footer)
    
    doc.build(elements)
    
    pdf = buffer.getvalue()
    buffer.close()
    response.write(pdf)
    
    return response


def export_co_attainment_excel(co_attainments, filters=None):
    """
    Generate an Excel report for Course Outcome attainment.
    
    Args:
        co_attainments: QuerySet of CourseOutcomeAttainment objects
        filters: Dictionary of applied filters (optional)
        
    Returns:
        HttpResponse: Excel file response
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"CO_Attainment_Report_{datetime.now().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "CO Attainment"
    
    # Define styles
    header_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='604993', end_color='604993', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    
    cell_alignment = Alignment(horizontal='left', vertical='center')
    center_alignment = Alignment(horizontal='center', vertical='center')
    
    border_style = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Add header row
    headers = ['Course Code', 'Course Name', 'CO Code', 'CO Description', 'Academic Year', 'Attainment %']
    for col_num, header in enumerate(headers, 1):
        cell = worksheet.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border_style
    
    # Add data rows
    for row_num, attainment in enumerate(co_attainments, 2):
        worksheet.cell(row=row_num, column=1).value = attainment.course_outcome.course.code
        worksheet.cell(row=row_num, column=2).value = attainment.course_outcome.course.name
        worksheet.cell(row=row_num, column=3).value = attainment.course_outcome.code
        worksheet.cell(row=row_num, column=4).value = attainment.course_outcome.description
        worksheet.cell(row=row_num, column=5).value = f"{attainment.academic_year.start_date.year}-{attainment.academic_year.end_date.year}"
        worksheet.cell(row=row_num, column=6).value = float(attainment.attainment_percentage) if attainment.attainment_percentage else 0
        
        # Apply styles to data cells
        for col_num in range(1, 7):
            cell = worksheet.cell(row=row_num, column=col_num)
            cell.border = border_style
            if col_num in [1, 3, 5, 6]:
                cell.alignment = center_alignment
            else:
                cell.alignment = cell_alignment
    
    # Adjust column widths
    worksheet.column_dimensions['A'].width = 15
    worksheet.column_dimensions['B'].width = 35
    worksheet.column_dimensions['C'].width = 12
    worksheet.column_dimensions['D'].width = 50
    worksheet.column_dimensions['E'].width = 18
    worksheet.column_dimensions['F'].width = 15
    
    # Save workbook to response
    workbook.save(response)
    
    return response


def export_po_attainment_excel(po_attainments, filters=None):
    """
    Generate an Excel report for Program Outcome attainment.
    
    Args:
        po_attainments: QuerySet of ProgramOutcomeAttainment objects
        filters: Dictionary of applied filters (optional)
        
    Returns:
        HttpResponse: Excel file response
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"PO_Attainment_Report_{datetime.now().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "PO Attainment"
    
    # Define styles
    header_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='604993', end_color='604993', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    
    cell_alignment = Alignment(horizontal='left', vertical='center')
    center_alignment = Alignment(horizontal='center', vertical='center')
    
    border_style = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Add header row
    headers = ['Department', 'PO Code', 'PO Description', 'Academic Year', 'Attainment %']
    for col_num, header in enumerate(headers, 1):
        cell = worksheet.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border_style
    
    # Add data rows
    for row_num, attainment in enumerate(po_attainments, 2):
        worksheet.cell(row=row_num, column=1).value = attainment.program_outcome.department.name
        worksheet.cell(row=row_num, column=2).value = attainment.program_outcome.code
        worksheet.cell(row=row_num, column=3).value = attainment.program_outcome.description
        worksheet.cell(row=row_num, column=4).value = f"{attainment.academic_year.start_date.year}-{attainment.academic_year.end_date.year}"
        worksheet.cell(row=row_num, column=5).value = float(attainment.attainment_percentage) if attainment.attainment_percentage else 0
        
        # Apply styles
        for col_num in range(1, 6):
            cell = worksheet.cell(row=row_num, column=col_num)
            cell.border = border_style
            if col_num in [2, 4, 5]:
                cell.alignment = center_alignment
            else:
                cell.alignment = cell_alignment
    
    # Adjust column widths
    worksheet.column_dimensions['A'].width = 25
    worksheet.column_dimensions['B'].width = 12
    worksheet.column_dimensions['C'].width = 55
    worksheet.column_dimensions['D'].width = 18
    worksheet.column_dimensions['E'].width = 15
    
    workbook.save(response)
    
    return response
